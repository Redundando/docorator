import io
from typing import Optional

import mammoth
import markdown
from cacherator import JSONCache, Cached
from docx import Document
from docx.document import Document as DocxDocument
from googleapiclient.http import MediaIoBaseUpload, MediaIoBaseDownload
from html2docx import html2docx
from logorator import Logger

from .auth import GoogleAuth
from .exceptions import DocumentNotFoundError, DocumentCreationError, DocumentSaveError


class Docorator(JSONCache):
    def __init__(self, keyfile_path: str, email: Optional[str], document_name: str, clear_cache: bool = True):
        super().__init__(data_id=f"{document_name}", directory="data/docorator", clear_cache=clear_cache)

        self.keyfile_path = keyfile_path
        self.email = email
        self.document_name = document_name

        self.auth = GoogleAuth(keyfile_path)
        self.docs_service = self.auth.get_docs_service()
        self.drive_service = self.auth.get_drive_service()

        if not hasattr(self, "document_id"):
            self.document_id = None

    def __str__(self):
        return f"{self.document_name} ({self.url})"

    def __repr__(self):
        return self.__str__()

    @property
    def url(self):
        if self.document_id:
            return f"https://docs.google.com/{self.document_id}"
        return None

    @Logger()
    def _find_document(self) -> bool:
        query = f"name = '{self.document_name}' and mimeType = 'application/vnd.google-apps.document'"
        results = self.drive_service.files().list(q=query, fields="files(id, name)").execute()
        files = results.get('files', [])

        if not files:
            return False

        self.document_id = files[0]['id']
        self.json_cache_save()
        return True

    @Logger()
    def _create_document(self) -> None:
        try:
            file_metadata = {
                'name'    : self.document_name,
                'mimeType': 'application/vnd.google-apps.document'
            }
            document = self.drive_service.files().create(body=file_metadata).execute()
            self.document_id = document['id']

            anyone_permission = {
                'type': 'anyone',
                'role': 'writer'
            }
            self.drive_service.permissions().create(
                fileId=self.document_id,
                body=anyone_permission
            ).execute()

            if self.email:
                try:
                    email_permission = {
                        'type'        : 'user',
                        'role'        : 'writer',
                        'emailAddress': self.email
                    }
                    self.drive_service.permissions().create(
                        fileId=self.document_id,
                        body=email_permission,
                        sendNotificationEmail=False,
                        fields='id'
                    ).execute()
                except Exception as e:
                    Logger.note(f"Warning: Failed to share document with {self.email}: {str(e)}")

            self.json_cache_save()
        except Exception as e:
            raise DocumentCreationError(f"Failed to create document: {str(e)}")

    @Logger()
    def _export_to_docx(self) -> DocxDocument:
        try:
            if not self.document_id:
                raise DocumentNotFoundError("Document ID not found")

            request = self.drive_service.files().export_media(
                fileId=self.document_id,
                mimeType='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )

            file = io.BytesIO()
            downloader = MediaIoBaseDownload(file, request)

            done = False
            while not done:
                status, done = downloader.next_chunk()

            file.seek(0)
            return Document(file)
        except Exception as e:
            raise DocumentNotFoundError(f"Failed to export document to DOCX: {str(e)}")

    @Cached()
    @Logger()
    def load(self) -> DocxDocument:
        if self.document_id:
            try:
                return self._export_to_docx()
            except DocumentNotFoundError:
                self.document_id = None

        if self._find_document():
            return self._export_to_docx()

        self._create_document()
        return Document()

    @Cached()
    @Logger()
    def as_markdown(self):
        buffer = io.BytesIO()
        doc = self.load()
        doc.save(buffer)
        buffer.seek(0)
        result = mammoth.convert_to_markdown(buffer)
        return result.value

    @Logger()
    def _convert_markdown_to_html(self, md: str = ""):
        html = markdown.markdown(
            md,
            extensions=[
                'markdown.extensions.tables',
                'markdown.extensions.fenced_code',
                'markdown.extensions.codehilite',
                'markdown.extensions.nl2br',
                'markdown.extensions.smarty',
                'markdown.extensions.toc'
            ]
        )

        full_html = f"""
                <!DOCTYPE html>
                <html>
                <body>
                    {html}
                </body>
                </html>
                """
        return full_html

    @Logger()
    def _convert_markdown_to_docx(self, md: str = ""):
        docx_bytes = html2docx(self._convert_markdown_to_html(md=md), title=self.document_name)
        return Document(docx_bytes)

    @Logger()
    def save(self, document: str | DocxDocument = "") -> bool:
        if isinstance(document, str):
            document = self._convert_markdown_to_docx(document)
        try:
            file_buffer = io.BytesIO()
            document.save(file_buffer)
            file_buffer.seek(0)

            media = MediaIoBaseUpload(
                file_buffer,
                mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                resumable=True
            )

            self.drive_service.files().update(
                fileId=self.document_id,
                media_body=media
            ).execute()
            return True
        except Exception as e:
            raise DocumentSaveError(f"Failed to save document: {str(e)}")
